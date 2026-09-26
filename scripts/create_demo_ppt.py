"""Generate the formal five-slide BeatIT demo presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "PPT" / "BeatIT_5_Slide_Demo.pptx"

NAVY = RGBColor(7, 20, 35)
NAVY_2 = RGBColor(12, 31, 51)
CYAN = RGBColor(45, 212, 191)
BLUE = RGBColor(56, 189, 248)
WHITE = RGBColor(247, 250, 252)
MUTED = RGBColor(166, 185, 204)
LINE = RGBColor(38, 63, 84)
RED = RGBColor(251, 113, 133)


def box(slide, x, y, w, h, fill=NAVY_2, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return frame


def title(slide, kicker, heading, subheading=None):
    text(slide, kicker.upper(), 0.7, 0.42, 11.9, 0.3, 10, CYAN, True)
    text(slide, heading, 0.7, 0.78, 11.9, 0.72, 28, WHITE, True)
    if subheading:
        text(slide, subheading, 0.72, 1.48, 11.5, 0.55, 13, MUTED)


def footer(slide, number):
    text(slide, "BEATIT · EDUCATIONAL CARDIAC SIMULATION · NOT A MEDICAL DEVICE",
         0.7, 7.12, 10.8, 0.2, 8, MUTED)
    text(slide, f"{number}/5", 12.0, 7.08, 0.6, 0.22, 9, CYAN, True, align=PP_ALIGN.RIGHT)


def metric(slide, value, label, x, y, w=2.3):
    box(slide, x, y, w, 1.05)
    text(slide, value, x + 0.18, y + 0.15, w - 0.36, 0.4, 24, CYAN, True)
    text(slide, label, x + 0.18, y + 0.6, w - 0.36, 0.25, 9, MUTED, True)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "BeatIT — Five-Slide Demo"
    prs.core_properties.subject = "Auditable multi-agent cardiac digital twin"
    prs.core_properties.author = "BeatIT"

    # 1 — opening
    slide = new_slide(prs)
    text(slide, "BEATIT", 0.75, 0.72, 4.5, 0.5, 14, CYAN, True)
    text(slide, "A cardiac digital twin\nthat knows what it does not know.", 0.75, 1.32, 7.4, 1.65, 34, WHITE, True)
    text(slide, "Real evidence in. Deterministic physiology. Explicit uncertainty out.",
         0.78, 3.15, 6.7, 0.6, 16, MUTED)
    box(slide, 8.25, 0.72, 4.3, 5.9, fill=NAVY_2)
    text(slide, "LIVE DEMO CONTRACT", 8.65, 1.08, 3.5, 0.3, 11, BLUE, True)
    for i, (label, value, color) in enumerate((
        ("REAL SUBJECT", "One dataset · one person", WHITE),
        ("OBSERVED", "Never inferred or stitched", CYAN),
        ("MISSING", "Shown explicitly", RED),
        ("MODEL PRIOR", "Labeled, bounded, inspectable", BLUE),
    )):
        y = 1.65 + i * 1.02
        text(slide, label, 8.65, y, 1.2, 0.25, 9, MUTED, True)
        text(slide, value, 9.85, y - 0.03, 2.2, 0.48, 14, color, True)
    text(slide, "The strongest claim is restraint.", 0.78, 5.75, 6.8, 0.5, 20, CYAN, True)
    footer(slide, 1)

    # 2 — problem
    slide = new_slide(prs)
    title(slide, "The clinical communication gap", "Cardiac evidence is fragmented. Explanations should not be.")
    for i, (name, body) in enumerate((
        ("EVIDENCE", "ECG, echo, vitals, reports and labs arrive through different workflows."),
        ("REASONING", "Clinicians must connect electrical, structural and hemodynamic signals under time pressure."),
        ("TRUST", "Most AI interfaces blur source evidence, derived values and model assumptions."),
    )):
        x = 0.75 + i * 4.15
        box(slide, x, 2.25, 3.75, 2.7)
        text(slide, f"0{i+1}", x + 0.25, 2.5, 0.6, 0.35, 16, CYAN, True)
        text(slide, name, x + 0.25, 3.02, 3.1, 0.35, 15, WHITE, True)
        text(slide, body, x + 0.25, 3.55, 3.12, 1.0, 12, MUTED)
    box(slide, 0.75, 5.3, 12.0, 1.05, fill=RGBColor(10, 40, 54), line=CYAN)
    text(slide, "BeatIT turns fragmented evidence into an inspectable physiological narrative—without crossing into diagnosis or treatment.",
         1.05, 5.57, 11.4, 0.45, 16, WHITE, True, align=PP_ALIGN.CENTER)
    footer(slide, 2)

    # 3 — workflow
    slide = new_slide(prs)
    title(slide, "Product", "One auditable pipeline. Eight specialist agents. Deterministic math.")
    steps = [
        ("01", "INGEST", "Real evidence"),
        ("02", "VALIDATE", "Source + quality"),
        ("03", "BUILD", "Twin state"),
        ("04", "SIMULATE", "Bounded scenarios"),
        ("05", "EXPLAIN", "Evidence + uncertainty"),
    ]
    for i, (n, label, detail) in enumerate(steps):
        x = 0.72 + i * 2.5
        box(slide, x, 2.15, 2.12, 1.48)
        text(slide, n, x + 0.18, 2.35, 0.4, 0.3, 10, CYAN, True)
        text(slide, label, x + 0.18, 2.75, 1.75, 0.28, 13, WHITE, True)
        text(slide, detail, x + 0.18, 3.12, 1.75, 0.22, 9, MUTED)
        if i < 4:
            text(slide, "→", x + 2.15, 2.67, 0.35, 0.35, 18, BLUE, True, align=PP_ALIGN.CENTER)
    metric(slide, "8", "SPECIALIST AGENTS", 0.75, 4.25)
    metric(slide, "1,331", "AUTOMATED TESTS", 3.25, 4.25)
    metric(slide, "0", "SYNTHETIC MEASUREMENTS", 5.75, 4.25)
    metric(slide, "0", "CROSS-PATIENT STITCHES", 8.25, 4.25)
    metric(slide, "100%", "SAFETY GATE", 10.75, 4.25, 2.0)
    text(slide, "Copilot orchestrates · Weave traces · Redis remembers · the physics core computes",
         1.0, 5.75, 11.3, 0.42, 17, CYAN, True, align=PP_ALIGN.CENTER)
    footer(slide, 3)

    # 4 — live demo
    slide = new_slide(prs)
    title(slide, "Live demonstration", "Two independent real-data heroes. Missing evidence remains missing.")
    box(slide, 0.75, 2.05, 5.75, 3.95)
    text(slide, "ECG HERO", 1.05, 2.35, 2.0, 0.3, 11, CYAN, True)
    text(slide, "PTB-XL", 1.05, 2.78, 2.2, 0.4, 22, WHITE, True)
    text(slide, "Real 12-lead ECG · 1,000 samples · diagnostic metadata", 1.05, 3.3, 4.9, 0.55, 12, MUTED)
    text(slide, "OBSERVED  ECG", 1.05, 4.05, 2.0, 0.3, 11, CYAN, True)
    text(slide, "MISSING    Echo · BP · medications", 1.05, 4.55, 4.6, 0.3, 11, RED, True)
    text(slide, "Demo moment: BeatIT widens uncertainty instead of inventing structure.", 1.05, 5.12, 4.9, 0.5, 12, WHITE, True)
    box(slide, 6.82, 2.05, 5.75, 3.95)
    text(slide, "CLINICAL HERO", 7.12, 2.35, 2.2, 0.3, 11, BLUE, True)
    text(slide, "UCI Heart Failure", 7.12, 2.78, 4.2, 0.4, 22, WHITE, True)
    text(slide, "Real clinical profile · EF · labs · risk context", 7.12, 3.3, 4.9, 0.55, 12, MUTED)
    text(slide, "OBSERVED  Clinical profile", 7.12, 4.05, 2.8, 0.3, 11, CYAN, True)
    text(slide, "MISSING    Raw ECG · raw echo · BP · medications", 7.12, 4.55, 4.9, 0.3, 11, RED, True)
    text(slide, "Demo moment: evidence, derivations and priors remain visually separable.", 7.12, 5.12, 4.9, 0.5, 12, WHITE, True)
    footer(slide, 4)

    # 5 — close
    slide = new_slide(prs)
    title(slide, "Why BeatIT wins", "Scientific restraint becomes a visible product advantage.")
    for i, (headline, body) in enumerate((
        ("TRUST", "Every number carries provenance, confidence and evidence class."),
        ("ORCHESTRATION", "Eight specialists collaborate over one deterministic physiology core."),
        ("DEMO READINESS", "Real ECG and clinical cases load directly; the entire repository gate passes."),
    )):
        y = 2.0 + i * 1.15
        text(slide, headline, 0.82, y, 2.0, 0.3, 12, CYAN, True)
        text(slide, body, 2.8, y - 0.05, 6.4, 0.58, 15, WHITE, True)
    box(slide, 9.5, 1.75, 3.05, 3.85, fill=RGBColor(10, 40, 54), line=CYAN)
    text(slide, "THE CLOSE", 9.85, 2.1, 2.35, 0.3, 11, MUTED, True, align=PP_ALIGN.CENTER)
    text(slide, "“BeatIT does not hallucinate completeness.”", 9.83, 2.75, 2.4, 1.25, 22, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, "It shows clinicians exactly what is known, derived, assumed—and missing.",
         9.85, 4.35, 2.35, 0.72, 12, CYAN, True, align=PP_ALIGN.CENTER)
    text(slide, "LET THE EVIDENCE LEAD.", 0.82, 5.65, 8.2, 0.55, 25, BLUE, True)
    footer(slide, 5)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
